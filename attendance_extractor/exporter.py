"""
Write the extracted record to JSON, CSV and a structured Excel workbook.

The Excel output reproduces the sheet's two-level header (DATE, MATIN>ENTRE/
SORTIE, SOIR>ENTRE/SORTIE) with merged cells, borders and centred text.
"""

from __future__ import annotations

import csv
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import EXCEL_HEADER, TABLE_COLUMNS
from utils import ensure_dir, get_logger

logger = get_logger()

_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_HEADER_FONT = Font(bold=True)


def export_json(record: dict, path: str) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(record, fh, ensure_ascii=False, indent=2)
    logger.info("Wrote JSON  -> %s", path)


def export_csv(record: dict, path: str) -> None:
    """Flat CSV: metadata block, a blank line, then the attendance table."""
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        for key in ("animateur", "magasin", "societe", "marque", "mois"):
            writer.writerow([key, record.get(key, "")])
        writer.writerow([])
        writer.writerow(["Date", "Matin Entre", "Matin Sortie",
                         "Soir Entre", "Soir Sortie"])
        for row in record.get("attendance", []):
            writer.writerow([row.get(c, "") for c in TABLE_COLUMNS])
    logger.info("Wrote CSV   -> %s", path)


def _write_metadata_block(ws, record: dict) -> int:
    """Write the header metadata at the top; return the next free row index."""
    fields = [
        ("Animateur", record.get("animateur", "")),
        ("Magasin", record.get("magasin", "")),
        ("Société", record.get("societe", "")),
        ("Marque", record.get("marque", "")),
        ("Mois", record.get("mois", "")),
    ]
    r = 1
    for label, value in fields:
        cell_label = ws.cell(row=r, column=1, value=label)
        cell_label.font = _HEADER_FONT
        ws.cell(row=r, column=2, value=value)
        r += 1
    return r + 1  # one blank spacer row


def _write_table_header(ws, start_row: int) -> int:
    """Write the two-level merged header. Return first data row index."""
    top = start_row
    sub = start_row + 1

    col = 1
    # DATE spans both header rows.
    ws.merge_cells(start_row=top, start_column=col, end_row=sub, end_column=col)
    _style_header(ws.cell(row=top, column=col, value="DATE"))
    col += 1

    for group, subs in EXCEL_HEADER.items():
        if group == "DATE":
            continue
        if subs:
            ws.merge_cells(start_row=top, start_column=col,
                           end_row=top, end_column=col + len(subs) - 1)
            _style_header(ws.cell(row=top, column=col, value=group))
            for i, sub_name in enumerate(subs):
                _style_header(ws.cell(row=sub, column=col + i, value=sub_name))
            col += len(subs)
        else:
            ws.merge_cells(start_row=top, start_column=col, end_row=sub, end_column=col)
            _style_header(ws.cell(row=top, column=col, value=group))
            col += 1

    # Make sure the merged sub-cells of DATE are also styled/bordered.
    _style_header(ws.cell(row=sub, column=1))
    return sub + 1


def _style_header(cell) -> None:
    cell.alignment = _CENTER
    cell.border = _BORDER
    cell.fill = _HEADER_FILL
    cell.font = _HEADER_FONT


def _write_table_body(ws, record: dict, first_row: int) -> None:
    r = first_row
    for row in record.get("attendance", []):
        for c, col_name in enumerate(TABLE_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(col_name, ""))
            cell.alignment = _CENTER
            cell.border = _BORDER
        r += 1


def export_excel(record: dict, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pointage"

    table_start = _write_metadata_block(ws, record)
    first_data_row = _write_table_header(ws, table_start)
    _write_table_body(ws, record, first_data_row)

    # Reasonable column widths.
    widths = [12, 14, 14, 14, 14]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(path)
    logger.info("Wrote Excel -> %s", path)


def export_all(record: dict, output_dir: str, stem: str) -> dict:
    """Write all three formats for one document; return their paths."""
    ensure_dir(output_dir)
    paths = {
        "json": os.path.join(output_dir, f"{stem}.json"),
        "csv": os.path.join(output_dir, f"{stem}.csv"),
        "xlsx": os.path.join(output_dir, f"{stem}.xlsx"),
    }
    export_json(record, paths["json"])
    export_csv(record, paths["csv"])
    export_excel(record, paths["xlsx"])
    return paths
